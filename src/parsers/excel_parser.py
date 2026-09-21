from __future__ import annotations

import json
from pathlib import Path
from typing import Optional
import openpyxl
from src.models.document import ParserResult, ParserStatus
from src.parsers.base import BaseParser
from src.parsers.usability_validator import assess_text_usability


class ExcelParser(BaseParser):
    """Parser for Excel (.xlsx) attachments (Design Extension, FR-006C).

    Textual Representation:
    - `text`, `raw_text`, and `clean_text` represent a source-derived textual serialization
      flattening workbook rows into delimited strings.
    - Delimiter ' | ' is parser-generated serialization syntax and NOT verbatim source content.
    - Sheet headers are excluded from `text` to prevent artificial strings from polluting source text.
    - Worksheet and cell provenance is preserved in `metadata['cell_coordinates']` (recording
      sheet name, openpyxl cell coordinate e.g. 'A1', row, column, and cell value).
    - Downstream evidence grounding for table cells must validate against actual cell values and
      metadata coordinates, NOT against serialization delimiters.
    """

    def parse(self, file_path: Path, document_id: Optional[str] = None) -> ParserResult:
        doc_id = document_id or file_path.name
        if not file_path.exists():
            return ParserResult(
                document_id=doc_id,
                status=ParserStatus.UNREADABLE,
                text=None,
                usable_for_extraction=False,
                diagnostic_evidence_ids=[f"missing_file_{doc_id}"],
                attempt_ids=[f"att_{doc_id}_xlsx"],
                error_message=f"File not found: {file_path}",
                metadata={"file_path": str(file_path)},
            )

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines: list[str] = []
            cell_coordinates: list[dict[str, int | str]] = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]

                for row in ws.iter_rows():
                    row_strs: list[str] = []
                    has_content = False

                    for cell in row:
                        val_str = str(cell.value).strip() if cell.value is not None else ""
                        row_strs.append(val_str)
                        if val_str:
                            has_content = True
                            cell_coordinates.append({
                                "sheet": sheet_name,
                                "coordinate": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "value": val_str,
                            })

                    if has_content:
                        lines.append(" | ".join(row_strs))

            full_text = "\n".join(lines).strip()

            assessment = assess_text_usability(full_text, is_pdf=False)

            diag_ids = [f"diag_{doc_id}_{r}" for r in assessment.reasons] if assessment.reasons else (
                [f"diag_{doc_id}_unreadable"] if assessment.status == ParserStatus.UNREADABLE else []
            )

            return ParserResult(
                document_id=doc_id,
                status=assessment.status,
                text=assessment.clean_text if assessment.is_usable else None,
                usable_for_extraction=assessment.is_usable,
                diagnostic_evidence_ids=diag_ids,
                attempt_ids=[f"att_{doc_id}_xlsx"],
                raw_text=full_text,
                clean_text=assessment.clean_text,
                is_scanned=False,
                error_message=None if assessment.is_usable else "; ".join(assessment.reasons),
                metadata={
                    "sheets": ",".join(wb.sheetnames),
                    "sheet_count": str(len(wb.sheetnames)),
                    "cell_coordinates": json.dumps(cell_coordinates),
                    "garbage_count": str(assessment.garbage_count),
                    "garbage_ratio": f"{assessment.garbage_ratio:.4f}",
                },
            )

        except Exception as e:
            return ParserResult(
                document_id=doc_id,
                status=ParserStatus.UNREADABLE,
                text=None,
                usable_for_extraction=False,
                diagnostic_evidence_ids=[f"diag_{doc_id}_corrupt_xlsx"],
                attempt_ids=[f"att_{doc_id}_xlsx"],
                is_scanned=False,
                error_message=f"Failed to parse XLSX: {type(e).__name__}: {e}",
                metadata={"error_type": type(e).__name__},
            )
